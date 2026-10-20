#!/usr/bin/env python3
"""Supplementary Table 6 in a portrait layout (A4, Word), for copying into the supplement.
One block per task: sample sizes in one merged cell, the three encoders as rows, each metric as
'value of the ten-model combination (mean of the ten models, range)'. Reads tables/SupplementaryTable6.csv."""
import os
import re

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import cswsi_config as cfg  # noqa: E402
T = cfg.TABLES
d = pd.read_csv(f'{T}/SupplementaryTable6.csv')


def comb(v, rng):
    """0.912 and '0.887 (0.850-0.916)' -> '0.91 (0.89, 0.85–0.92)'."""
    if not isinstance(rng, str) or not re.match(r'^[0-9.]+ \(', rng):
        return str(v) if pd.notna(v) else '–'
    m = re.match(r'([0-9.]+) \(([0-9.]+)-([0-9.]+)\)', rng)
    mean, lo, hi = (float(x) for x in m.groups())
    return f'{float(v):.2f} ({mean:.2f}, {lo:.2f}–{hi:.2f})'


def clean(x):
    return '–' if (not isinstance(x, str) and pd.isna(x)) or x == 'not evaluated' else str(x)


doc = Document()
s = doc.sections[0]
s.page_height, s.page_width = Cm(29.7), Cm(21.0)
s.left_margin = s.right_margin = Cm(1.6)
s.top_margin = s.bottom_margin = Cm(1.6)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(7.5)
st.paragraph_format.space_after = Pt(0)
p = doc.add_paragraph()
r = p.add_run('Supplementary Table 6. Whole-slide image benchmark per task and encoder. ')
r.bold = True; r.font.size = Pt(9)
r = p.add_run('Held-out evaluation set of the discovery cohort (30% of patients) and, where evaluated, CPTAC-COAD '
              '(20×, StainNet-normalized). AUC: of the probability averaged over the ten cross-validation models; '
              'balanced accuracy: of their majority vote. In parentheses: mean of the ten models, and their range.')
r.font.size = Pt(8)
p.paragraph_format.space_after = Pt(6)

hdr = ['Task', 'Slides (patients); slides per class', 'Encoder', 'Held-out AUC', 'Held-out balanced accuracy',
       'CPTAC AUC', 'CPTAC balanced accuracy (majority vote)']
widths = [Cm(3.3), Cm(4.3), Cm(1.5), Cm(2.4), Cm(2.4), Cm(1.9), Cm(2.0)]
t = doc.add_table(rows=1, cols=len(hdr)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False


def shade(cell, color):
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(sh)


def put(cell, text, bold=False, size=7.5):
    cell.text = ''
    lines = str(text).split('\n')
    for k, line in enumerate(lines):
        par = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        rr = par.add_run(line); rr.bold = bold; rr.font.size = Pt(size)


for i, h in enumerate(hdr):
    put(t.rows[0].cells[i], h, bold=True); shade(t.rows[0].cells[i], 'DCE6F0')
# repeat header row on every page
trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)

for task, g in d.groupby('Task', sort=False):
    first = g.iloc[0]
    n = f"Held-out: {first['Held-out: slides (patients)']}\n{first['Held-out: slides per class']}"
    cp = clean(first['CPTAC 20x StainNet: slides (cases)'])
    if cp != '–':
        n += f"\nCPTAC: {cp}\n{first['CPTAC 20x StainNet: slides per class']}"
    rows = []
    for _, r in g.iterrows():
        cells = t.add_row().cells
        put(cells[2], r.Encoder)
        put(cells[3], comb(r['Held-out: AUC, averaged prediction'], r['Held-out: AUC, mean of the ten models (range)']))
        put(cells[4], comb(r['Held-out: balanced accuracy, majority vote'],
                           r['Held-out: balanced accuracy, mean of the ten models (range)']))
        ca = r['CPTAC 20x StainNet: AUC, averaged prediction']
        put(cells[5], comb(ca, r['CPTAC 20x StainNet: AUC, mean of the ten models (range)']) if clean(ca) != '–' and not str(ca).startswith('n/a') else ('n/a' if str(ca).startswith('n/a') else '–'))
        cb = clean(r['CPTAC 20x StainNet: balanced accuracy, majority vote'])
        if cb.startswith('sensitivity'):
            cb = f"{float(cb.split()[-1]):.1%} called tumor"
        elif cb != '–':
            cb = f'{float(cb):.2f}'
        put(cells[6], cb)
        rows.append(cells)
    a, b = rows[0], rows[-1]
    a[0].merge(b[0]); a[1].merge(b[1])
    put(a[0], task, bold=True); put(a[1], n, size=7)
for row in t.rows:
    for i, c in enumerate(row.cells):
        c.width = widths[i]
doc.add_paragraph()
notes = pd.read_excel(f'{T}/SupplementaryTable6.xlsx', sheet_name='Notes').Note.tolist()
for nt in notes:
    q = doc.add_paragraph(nt); q.runs[0].font.size = Pt(7)
    q.paragraph_format.space_after = Pt(2)
doc.save(f'{T}/SupplementaryTable6_portrait.docx')
print('wrote', f'{T}/SupplementaryTable6_portrait.docx')
