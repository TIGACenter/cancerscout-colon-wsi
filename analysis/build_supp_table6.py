#!/usr/bin/env python3
"""
Supplementary Table 6: whole-slide-image benchmark metrics per task and encoder, on the held-out
evaluation set and, for the endpoints evaluated there, on CPTAC (matched 20x, StainNet).
Reads tables/in_domain_metrics.csv, in_domain_perfold.csv, cptac_metrics.csv, cptac_perfold.csv
(compute_numbers.py). Writes tables/SupplementaryTable6.{csv,xlsx}.
"""
import os

import warnings

import numpy as np
import pandas as pd

warnings.filterwarnings('ignore', category=pd.errors.PerformanceWarning)

import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import cswsi_config as cfg  # noqa: E402
T = cfg.TABLES
ENC = {'resnet': 'ResNet50', 'uni': 'UNI', 'conch': 'CONCH'}
TASKS = [('tumor_vs_normal', 'Tumor detection (tumor vs tumor-free)'), ('sex', 'Sex'), ('grading', 'Grade (3-class: G2, G3, MSI tumors without a grade)'),
         ('ln_metastasis', 'Lymph-node metastasis'), ('distant_metastasis', 'Distant metastasis'),
         ('cptac', 'CPTAC subtype (CIN / MSI / mesenchymal)'), ('msi_status', 'MSI status (IHC label)'),
         ('msi_status_genetic', 'MSI status (genetic label)'), ('msi_status_consensus', 'MSI status (consensus label)'),
         ('braf', 'BRAF mutation'), ('k_n_ras', 'KRAS/NRAS mutation'), ('tp53', 'TP53 mutation'), ('apc', 'APC mutation'),
         ('proteomic_classification', 'Proteomic subgroups C1-C6 (6-class)'),
         ('p_g_c2c3c5_cin_cases', 'Proteomic subgroups C2/C3/C5, CPTAC-CIN cases (3-class)'),
         ('cin_vs_mcml', 'CIN vs mesenchymal')]


EVALD = cfg.EVAL_DIR
SUBNAMES = {'1': 'C1', '2': 'C2', '3': 'C3', '4': 'C4', '5': 'C5', '6': 'C6'}


def class_counts(d, task):
    """'name n; name n' on the slides of one evaluation (fold 0; every fold scores the same slides)."""
    import glob as _g, json as _j
    f = _g.glob(f'{d}/fold_0.csv')
    if not f:
        return ''
    y = pd.read_csv(f[0]).drop_duplicates('slide_id').Y.astype(int)
    inv = {v: k for k, v in _j.load(open(f'{d}/label_dict.json')).items()}
    names = {k: (SUBNAMES.get(str(inv[k]), str(inv[k])) if 'proteomic' in task or 'p_g_' in task else str(inv[k]))
             for k in inv}
    if task == 'grading':
        names = {k: {'2': 'G2', '3': 'G3', 'MSI': 'MSI (no grade given)'}.get(str(v), str(v)) for k, v in names.items()}
    c = y.value_counts()
    return '; '.join(f'{names[k]} {int(n)}' for k, n in sorted(c.items(), key=lambda kv: names[kv[0]]))


def fmt_range(pf, key, col):
    v = pf.loc[key, col].dropna()
    return f'{v.mean():.3f} ({v.min():.3f}-{v.max():.3f})'


ind = pd.read_csv(f'{T}/in_domain_metrics.csv').set_index(['task', 'encoder'])
ipf = pd.read_csv(f'{T}/in_domain_perfold.csv').set_index(['task', 'encoder'])
cp = pd.read_csv(f'{T}/cptac_metrics.csv')
cp = cp[cp.condition == '20x StainNet'].set_index(['task', 'encoder'])
cpf = pd.read_csv(f'{T}/cptac_perfold.csv')
cpf = cpf[cpf.condition == '20x StainNet'].set_index(['task', 'encoder'])
rows = []
for t, name in TASKS:
    for e, en in ENC.items():
        r = ind.loc[(t, e)]
        row = {'Task': name, 'Encoder': en,
               'Held-out: slides (patients)': f'{int(r.n_slides)} ({int(r.n_patients)})',
               'Held-out: slides per class': class_counts(f'{EVALD}/EVAL_cancerscout_colon_he_{t}_v3_{e}_s1', t),
               'Held-out: AUC, averaged prediction': round(r.auc_consensus, 3),
               'Held-out: AUC, mean of the ten models (range)': fmt_range(ipf, (t, e), 'auc'),
               'Held-out: balanced accuracy, majority vote': round(r.bacc_consensus, 3),
               'Held-out: balanced accuracy, mean of the ten models (range)': fmt_range(ipf, (t, e), 'bacc')}
        if (t, e) in cp.index:
            c = cp.loc[(t, e)]
            row['CPTAC 20x StainNet: slides (cases)'] = f'{int(c.n_slides)} ({int(c.n_patients)})'
            row['CPTAC 20x StainNet: slides per class'] = class_counts(f'{EVALD}/EVAL_cancerscout_colon-cptac_{t}_v3_{e}_scale20xsn_full', t)
            if t == 'tumor_vs_normal':
                row['CPTAC 20x StainNet: AUC, averaged prediction'] = 'n/a (all slides tumor)'
                row['CPTAC 20x StainNet: AUC, mean of the ten models (range)'] = 'n/a'
                row['CPTAC 20x StainNet: balanced accuracy, majority vote'] = (
                    f'sensitivity {c.called_class0_vote:.3f}')
            else:
                row['CPTAC 20x StainNet: AUC, averaged prediction'] = round(c.auc_consensus, 3)
                row['CPTAC 20x StainNet: AUC, mean of the ten models (range)'] = fmt_range(cpf, (t, e), 'auc')
                row['CPTAC 20x StainNet: balanced accuracy, majority vote'] = round(c.bacc_consensus, 3)
        else:
            for k in ('slides (cases)', 'slides per class', 'AUC, averaged prediction', 'AUC, mean of the ten models (range)', 'balanced accuracy, majority vote'):
                row[f'CPTAC 20x StainNet: {k}'] = 'not evaluated'
        rows.append(row)
d = pd.DataFrame(rows)
d.to_csv(f'{T}/SupplementaryTable6.csv', index=False)
notes = pd.DataFrame({'Note': [
    'Held-out set: the patient-level evaluation split (30% of patients), one row per slide; the ten cross-validation models were trained on the remaining 70%. Slides per class: slides of the evaluation set in each class of the task (identical for the three encoders).',
    'AUC, averaged prediction: AUC of the predicted probability averaged over the ten cross-validation models (this is not the mean of the ten AUCs, which is given in the next column). Multiclass AUC: macro-averaged one-vs-rest.',
    'Balanced accuracy: each model called at the threshold fitted on its own internal test split of the training set (binary) or by argmax (multiclass); majority vote = the class called by most of the ten models, ties broken by the averaged probability.',
    'CPTAC: CPTAC-COAD slides (TCIA) tiled at the training resolution (20x, ~0.5 um/px, 512 px level-0 crops) and stain-normalized with StainNet; thresholds transferred from CancerScout, not recalibrated.',
    'MSI on CPTAC: the three CancerScout models (IHC-, genetic- and consensus-trained) are scored against the single CPTAC MSI label (223 slides, 53 MSI).',
    'Grade: the grading labels give MSI tumors no grade but the value MSI, so the task has three classes (G2, G3, MSI) and its AUC partly reflects MSI detection.',
    'Tumor detection on CPTAC: all slides are tumor-positive, so only the share called tumor by the majority vote (sensitivity) is defined.',
    'CIN subgroups: the three CIN subgroups C2, C3 and C5, on slides of CPTAC-CIN tumors only (same training method, patient split and ten-model design as every other task).']})
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
wb = Workbook()
ws = wb.active
ws.title = 'Supplementary Table 6'
ws.cell(1, 1, 'Supplementary Table 6. Whole-slide image benchmark per task and encoder: held-out evaluation set of '
              'the discovery cohort and, where evaluated, CPTAC-COAD (20x, StainNet).').font = Font(bold=True, size=12)
cols = list(d.columns)
groups = [('Task', ['Task', 'Encoder'])]
ho = [c for c in cols if c.startswith('Held-out: ')]
cp = [c for c in cols if c.startswith('CPTAC 20x StainNet: ')]
groups += [('Held-out evaluation set (discovery cohort, 30% of patients)', ho),
           ('CPTAC-COAD, 20x, StainNet-normalized (external)', cp)]
fill = PatternFill('solid', fgColor='DCE6F0')
thin = Side(style='thin', color='B0B8C4')
col = 1
for title, members in groups:
    ws.cell(3, col, title if title != 'Task' else '')
    if len(members) > 1:
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + len(members) - 1)
    for m in members:
        c = ws.cell(4, col, m.replace('Held-out: ', '').replace('CPTAC 20x StainNet: ', ''))
        c.font = Font(bold=True); c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = Border(bottom=thin)
        col += 1
for c in ws[3]:
    c.font = Font(bold=True); c.alignment = Alignment(horizontal='center')
order = [m for _, ms in groups for m in ms]
for i, r in enumerate(d[order].itertuples(index=False), start=5):
    for j, v in enumerate(r, start=1):
        ws.cell(i, j, v)
widths = {'Task': 46, 'Encoder': 10}
for j, m in enumerate(order, start=1):
    ws.column_dimensions[get_column_letter(j)].width = widths.get(m, 22 if 'per class' not in m else 30)
ws.freeze_panes = 'C5'
ws.row_dimensions[4].height = 45
wn = wb.create_sheet('Notes')
wn.cell(1, 1, 'Note').font = Font(bold=True)
for i, n in enumerate(notes.Note, start=2):
    wn.cell(i, 1, n)
wn.column_dimensions['A'].width = 160
wb.save(f'{T}/SupplementaryTable6.xlsx')
print(d.to_string(index=False, max_colwidth=40))
